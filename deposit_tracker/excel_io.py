"""Export to / import from the deposit-slip workbook layout (Summary + one sheet per account)."""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .storage import Database, Settings, Slip

HEADERS = [
    "#", "Deposit Date", "Time", "TD Ref #", "Occasion / Collection Date (per note)",
    "Handwritten Note", "Contributors", "CAD Cash ($)", "USD Cash (US$)",
    "USD Converted ($ CAD)", "FX Rate", "Total Deposited ($ CAD)", "Denominations",
    "Remarks", "Slip Image",
]
WIDTHS = [5, 13, 8, 20, 30, 45, 13, 13, 13, 16, 9, 18, 32, 40, 14]
FIRST_ROW = 5

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00'


def _sheet_title(name: str) -> str:
    bad = '[]:*?/\\'
    clean = "".join("-" if c in bad else c for c in name).strip()
    return clean[:31] or "Sheet"


def _to_date(value: str) -> date | str:
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        return value or None


def _to_time(value: str) -> time | str:
    try:
        return time.fromisoformat(value)
    except (TypeError, ValueError):
        return value or None


def export_workbook(path: str | Path, db: Database, settings: Settings,
                    slips: list[Slip] | None = None) -> None:
    slips = db.all_slips() if slips is None else slips
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    by_account: dict[str, list[Slip]] = {}
    for name in settings.account_names():
        by_account[name] = []
    for slip in slips:
        by_account.setdefault(slip.account or "Unassigned", []).append(slip)

    sheet_refs = []  # (account name, sheet title, total row)
    used_titles = {"Summary"}
    for account, rows in by_account.items():
        if not rows:
            continue
        cfg = next((a for a in settings.accounts if a["name"] == account), None)
        title = _sheet_title(cfg["sheet"] if cfg else account)
        while title in used_titles:
            title = _sheet_title(title[:28] + "-2")
        used_titles.add(title)
        ws = wb.create_sheet(title)
        total_row = _write_account_sheet(ws, account, rows, db, settings)
        sheet_refs.append((account, title, total_row))

    _write_summary(summary, sheet_refs, slips, settings)
    wb.save(path)


def _write_account_sheet(ws, account: str, rows: list[Slip], db: Database,
                         settings: Settings) -> int:
    ws["A1"] = f"{account} — Deposits to {settings.organization} ({settings.branch})"
    ws["A1"].font = TITLE
    ws["A2"] = ("Transcribed from bank deposit slips. Deposit date is the bank date; "
                "the collection date is from the handwritten note.")
    ws["A2"].font = Font(italic=True, color="595959")

    for col, (head, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = ws.cell(row=4, column=col, value=head)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[4].height = 32

    rows = sorted(rows, key=lambda s: (s.deposit_date, s.deposit_time))
    for i, slip in enumerate(rows):
        r = FIRST_ROW + i
        values = [
            i + 1, _to_date(slip.deposit_date), _to_time(slip.deposit_time), slip.ref_no,
            slip.occasion, slip.note, slip.contributors, slip.cad_cash, slip.usd_cash,
            slip.usd_converted, slip.fx_rate, f"=H{r}+J{r}", slip.denominations, slip.remarks,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=col in (5, 6, 13, 14), vertical="top")
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=3).number_format = "hh:mm"
        for col in (8, 9, 10, 12):
            ws.cell(row=r, column=col).number_format = MONEY
        ws.cell(row=r, column=11).number_format = "0.0000"

        img_cell = ws.cell(row=r, column=15)
        img_cell.border = BOX
        image = db.image_abspath(slip)
        if image:
            img_cell.value = "Open image"
            img_cell.hyperlink = image.resolve().as_uri()
            img_cell.font = Font(color="0563C1", underline="single")

    last = FIRST_ROW + len(rows) - 1
    total_row = last + 1
    ws.cell(row=total_row, column=1, value="TOTAL")
    for col in (7, 8, 9, 10, 12):
        letter = get_column_letter(col)
        ws.cell(row=total_row, column=col, value=f"=SUM({letter}{FIRST_ROW}:{letter}{last})")
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.font, cell.fill, cell.border = BOLD, TOTAL_FILL, BOX
        if col in (8, 9, 10, 12):
            cell.number_format = MONEY

    ws.cell(row=total_row + 2, column=1,
            value="Contributor counts are only where written on the slip; blank = not noted, "
                  "so the total is a partial count.").font = Font(italic=True, color="595959")
    ws.freeze_panes = "C5"
    return total_row


def _write_summary(ws, sheet_refs, slips: list[Slip], settings: Settings) -> None:
    dates = sorted(s.deposit_date for s in slips if s.deposit_date)
    period = ""
    if dates:
        try:
            first, last = date.fromisoformat(dates[0]), date.fromisoformat(dates[-1])
            period = f" Period: {first:%b %d, %Y} – {last:%b %d, %Y}."
        except ValueError:
            pass
    ws["A1"] = f"Deposit Slip Summary — {settings.organization}"
    ws["A1"].font = TITLE
    ws["A2"] = f"All deposits made at {settings.branch}.{period}"
    ws["A2"].font = Font(italic=True, color="595959")

    heads = ["Account", "# Slips", "CAD Cash ($)", "USD Converted ($ CAD)",
             "Total Deposited ($ CAD)", "Contributors Noted"]
    for col, (head, width) in enumerate(zip(heads, [30, 10, 15, 22, 24, 20]), start=1):
        cell = ws.cell(row=4, column=col, value=head)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BOX
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    r = 5
    for account, title, total_row in sheet_refs:
        q = f"'{title}'"
        last = total_row - 1
        values = [account, f"=COUNT({q}!A{FIRST_ROW}:A{last})", f"={q}!H{total_row}",
                  f"={q}!J{total_row}", f"={q}!L{total_row}", f"={q}!G{total_row}"]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = BOX
            if col in (3, 4, 5):
                cell.number_format = MONEY
        r += 1

    ws.cell(row=r, column=1, value="GRAND TOTAL")
    if sheet_refs:
        for col in range(2, 7):
            letter = get_column_letter(col)
            ws.cell(row=r, column=col, value=f"=SUM({letter}5:{letter}{r - 1})")
    for col in range(1, 7):
        cell = ws.cell(row=r, column=col)
        cell.font, cell.fill, cell.border = BOLD, TOTAL_FILL, BOX
        if col in (3, 4, 5):
            cell.number_format = MONEY

    remarks = [s for s in sorted(slips, key=lambda s: s.deposit_date) if s.remarks.strip()]
    if remarks:
        ws.cell(row=r + 2, column=1, value="Notes:").font = BOLD
        for i, slip in enumerate(remarks, start=r + 3):
            ws.cell(row=i, column=1,
                    value=f"• {slip.deposit_date} {slip.account} ({slip.ref_no}): {slip.remarks}")


# --- import ------------------------------------------------------------------

HEADER_MAP = {
    "deposit date": "deposit_date", "time": "deposit_time", "td ref #": "ref_no",
    "ref #": "ref_no", "reference": "ref_no",
    "occasion / collection date (per note)": "occasion", "occasion": "occasion",
    "handwritten note": "note", "contributors": "contributors",
    "cad cash ($)": "cad_cash", "usd cash (us$)": "usd_cash",
    "usd converted ($ cad)": "usd_converted", "fx rate": "fx_rate",
    "denominations": "denominations", "remarks": "remarks",
}


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value).strip()


def _cell_to_num(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).replace("$", "").replace(",", ""))
    except ValueError:
        return None


def import_workbook(path: str | Path, db: Database, settings: Settings) -> tuple[int, int]:
    """Import every account sheet. Returns (imported, skipped as duplicates)."""
    wb = load_workbook(path, data_only=True)
    imported = skipped = 0
    for ws in wb.worksheets:
        header_row = None
        for row in ws.iter_rows(min_row=1, max_row=15):
            if str(row[0].value).strip() == "#":
                header_row = row[0].row
                break
        if header_row is None:
            continue

        columns = {}
        for cell in ws[header_row]:
            key = HEADER_MAP.get(str(cell.value or "").strip().lower())
            if key:
                columns[cell.column - 1] = key
        if "deposit_date" not in columns.values():
            continue

        account = _account_for_sheet(ws, settings)
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not isinstance(row[0], (int, float)):
                continue  # TOTAL row, notes, blanks
            slip = Slip(account=account, branch=settings.branch)
            for idx, key in columns.items():
                value = row[idx] if idx < len(row) else None
                if key in ("cad_cash", "usd_cash", "usd_converted", "fx_rate"):
                    setattr(slip, key, _cell_to_num(value))
                elif key == "contributors":
                    num = _cell_to_num(value)
                    slip.contributors = int(num) if num is not None else None
                else:
                    setattr(slip, key, _cell_to_str(value))
            if db.find_by_ref(slip.ref_no):
                skipped += 1
                continue
            db.save(slip)
            imported += 1
    return imported, skipped


def _account_for_sheet(ws, settings: Settings) -> str:
    for acc in settings.accounts:
        if acc["sheet"].lower() == ws.title.lower() or acc["name"].lower() == ws.title.lower():
            return acc["name"]
    title = str(ws["A1"].value or "")
    name = title.split(" — ")[0].strip() if " — " in title else ws.title
    for acc in settings.accounts:
        if acc["name"].lower() == name.lower():
            return acc["name"]
    settings.accounts.append({"name": name, "sheet": ws.title, "keywords": ""})
    settings.save()
    return name
